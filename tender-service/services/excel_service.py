"""
Excel-backed tender storage.

This is a temporary persistence layer for fast product validation. The public
methods mirror the TenderStorage contract so a database implementation can
replace this class later with minimal application changes.
"""

from __future__ import annotations

import asyncio
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from models.tender import Tender
from services.tender_storage import TenderStorage
from utils.logger import get_logger

logger = get_logger(__name__)


RAW_TENDERS_SHEET = "Raw_Tenders"
AI_ANALYSIS_SHEET = "AI_Analysis"
PROCESSING_STATUS_SHEET = "Processing_Status"

LEGACY_SHEET_NAMES = {
    "Raw Tenders": RAW_TENDERS_SHEET,
    "AI Analysis": AI_ANALYSIS_SHEET,
    "Processing Status": PROCESSING_STATUS_SHEET,
}

RAW_TENDER_HEADERS = [
    "Reference",
    "Tender ID",
    "Title",
    "Organisation",
    "Department",
    "Tender Number",
    "Tender Category",
    "Tender Type",
    "Form Of Contract",
    "Withdrawal Allowed",
    "No. of Covers",
    "Payment Mode",
    "Published Date",
    "Document Download Start Date",
    "Document Download End Date",
    "Clarification Start Date",
    "Clarification End Date",
    "Bid Submission Start Date",
    "Bid Submission End Date",
    "Closing Date",
    "Opening Date",
    "Pre Bid Meeting Date",
    "Pre Bid Meeting Place",
    "Bid Opening Place",
    "Tender Value",
    "EMD",
    "EMD Exemption",
    "EMD Fee Type",
    "Tender Fee",
    "Fee Payable To",
    "Fee Payable At",
    "Description",
    "Product Category",
    "Sub category",
    "Contract Type",
    "Bid Validity",
    "Period Of Work",
    "Location",
    "Pincode",
    "Pre Qualification",
    "Inviting Authority Name",
    "Inviting Authority Address",
    "PDF URL",
    "PDF Path",
    "Supporting Documents",
]

AI_ANALYSIS_HEADERS = [
    "Reference",
    "Energy Related",
    "Energy Sector",
    "Headline",
    "Summary",
    "SEO Title",
    "Meta Description",
    "Slug",
    "Keywords",
    "Tags",
    "Importance Score",
    "Publish Recommendation",
    "Why It Matters",
    "Industry Impact",
    "Companies Interested",
    "Eligibility",
    "Required Documents",
    "Important Dates",
    "Risk Level",
]

PROCESSING_STATUS_HEADERS = [
    "Reference",
    "Downloaded",
    "PDF Parsed",
    "AI Completed",
    "Created At",
    "Updated At",
]

SHEET_HEADERS = {
    RAW_TENDERS_SHEET: RAW_TENDER_HEADERS,
    AI_ANALYSIS_SHEET: AI_ANALYSIS_HEADERS,
    PROCESSING_STATUS_SHEET: PROCESSING_STATUS_HEADERS,
}

AI_ANALYSIS_FIELD_ALIASES = {
    "is_energy_related": "Energy Related",
    "energy_related": "Energy Related",
    "energy_sector": "Energy Sector",
    "headline": "Headline",
    "summary": "Summary",
    "seo_title": "SEO Title",
    "meta_description": "Meta Description",
    "slug": "Slug",
    "keywords": "Keywords",
    "tags": "Tags",
    "importance_score": "Importance Score",
    "publish_recommendation": "Publish Recommendation",
    "bid_recommendation": "Publish Recommendation",
    "why_it_matters": "Why It Matters",
    "industry_impact": "Industry Impact",
    "companies_interested": "Companies Interested",
    "eligibility": "Eligibility",
    "required_documents": "Required Documents",
    "important_dates": "Important Dates",
    "risk_level": "Risk Level",
}

PROCESSING_STATUS_FIELD_ALIASES = {
    "downloaded": "Downloaded",
    "pdf_parsed": "PDF Parsed",
    "ai_completed": "AI Completed",
    "created_at": "Created At",
    "updated_at": "Updated At",
}


class ExcelService(TenderStorage):
    """Persist tender data in a multi-sheet Excel workbook."""

    def __init__(self, workbook_path: str | Path) -> None:
        self.workbook_path = Path(workbook_path)
        self._lock = asyncio.Lock()

    async def initialize(self) -> None:
        """Create the workbook and required sheets if they do not exist."""
        async with self._lock:
            await asyncio.to_thread(self._initialize_sync)

    async def tender_exists(self, reference: str) -> bool:
        """Check duplicate tenders by reference number."""
        async with self._lock:
            return await asyncio.to_thread(self._tender_exists_sync, reference)

    async def ai_analysis_completed(self, reference: str) -> bool:
        """Check whether AI analysis has already completed for a reference."""
        async with self._lock:
            return await asyncio.to_thread(
                self._ai_analysis_completed_sync,
                reference,
            )

    async def append_tender(self, tender: Tender) -> bool:
        """
        Append a tender to all baseline sheets.

        The AI row is created with blank fields so later analysis updates can
        target the same reference without needing to create another row.
        """
        async with self._lock:
            return await asyncio.to_thread(self._append_tender_sync, tender)

    async def update_ai_analysis(
        self,
        reference: str,
        analysis: Mapping[str, Any],
    ) -> None:
        """Create or update AI analysis fields for a tender reference."""
        async with self._lock:
            await asyncio.to_thread(
                self._update_row_sync,
                AI_ANALYSIS_SHEET,
                reference,
                analysis,
                AI_ANALYSIS_FIELD_ALIASES,
            )

    async def update_processing_status(
        self,
        reference: str,
        status: Mapping[str, Any],
    ) -> None:
        """Create or update processing status fields for a tender reference."""
        enriched_status = dict(status)
        enriched_status.setdefault("updated_at", self._now())

        async with self._lock:
            await asyncio.to_thread(
                self._update_row_sync,
                PROCESSING_STATUS_SHEET,
                reference,
                enriched_status,
                PROCESSING_STATUS_FIELD_ALIASES,
            )

    def _initialize_sync(self) -> None:
        self.workbook_path.parent.mkdir(parents=True, exist_ok=True)

        if self.workbook_path.exists():
            workbook = load_workbook(self.workbook_path)
        else:
            workbook = Workbook()
            default_sheet = workbook.active
            workbook.remove(default_sheet)

        self._rename_legacy_sheets(workbook)

        for sheet_name, headers in SHEET_HEADERS.items():
            worksheet = self._get_or_create_sheet(workbook, sheet_name)
            self._ensure_headers(worksheet, headers)
            self._format_header(worksheet)

        workbook.save(self.workbook_path)
        logger.info("Excel workbook ready: %s", self.workbook_path)

    def _save_with_retry(self, workbook, max_retries: int = 3) -> None:
        """Save workbook with retry logic for Windows file-lock conflicts."""
        for attempt in range(1, max_retries + 1):
            try:
                workbook.save(self.workbook_path)
                return
            except PermissionError:
                if attempt == max_retries:
                    logger.error(
                        "Cannot save Excel file: %s — "
                        "Close the file in Microsoft Excel and try again.",
                        self.workbook_path,
                    )
                    raise
                delay = attempt * 2
                logger.warning(
                    "Excel file is locked (attempt %d/%d). "
                    "Retrying in %ds... Close the file if open in Excel.",
                    attempt, max_retries, delay,
                )
                time.sleep(delay)

    def _tender_exists_sync(self, reference: str) -> bool:
        workbook = self._load_ready_workbook()
        worksheet = workbook[RAW_TENDERS_SHEET]
        return self._reference_exists(worksheet, reference)

    def _ai_analysis_completed_sync(self, reference: str) -> bool:
        workbook = self._load_ready_workbook()
        worksheet = workbook[PROCESSING_STATUS_SHEET]
        row_number = self._find_reference_row(worksheet, reference)

        if row_number is None:
            return False

        header_indexes = self._header_indexes(worksheet)
        column = header_indexes["AI Completed"]
        value = worksheet.cell(row=row_number, column=column).value
        return self._truthy(value)

    def _append_tender_sync(self, tender: Tender) -> bool:
        workbook = self._load_ready_workbook()
        raw_sheet = workbook[RAW_TENDERS_SHEET]

        if self._reference_exists(raw_sheet, tender.reference):
            logger.info("Duplicate tender skipped: %s", tender.reference)
            return False

        created_at = self._now()

        raw_sheet.append(self._raw_tender_row(tender))
        workbook[AI_ANALYSIS_SHEET].append(
            [tender.reference] + [None] * (len(AI_ANALYSIS_HEADERS) - 1)
        )
        workbook[PROCESSING_STATUS_SHEET].append(
            [
                tender.reference,
                bool(tender.pdf_local_path),
                bool(tender.extracted_text),
                False,
                created_at,
                created_at,
            ]
        )

        self._save_with_retry(workbook)
        logger.info("Tender stored in Excel: %s", tender.reference)
        return True

    def _update_row_sync(
        self,
        sheet_name: str,
        reference: str,
        values: Mapping[str, Any],
        aliases: Mapping[str, str],
    ) -> None:
        workbook = self._load_ready_workbook()
        worksheet = workbook[sheet_name]

        row_number = self._find_reference_row(worksheet, reference)
        if row_number is None:
            worksheet.append([reference] + [None] * (worksheet.max_column - 1))
            row_number = worksheet.max_row

        header_indexes = self._header_indexes(worksheet)

        for field_name, value in values.items():
            header = self._resolve_header(field_name, aliases)
            if header == "Reference":
                continue
            if header not in header_indexes:
                raise ValueError(f"Unsupported field for sheet '{sheet_name}': {field_name}")
            column = header_indexes[header]
            worksheet.cell(row=row_number, column=column, value=self._cell_value(value))

        self._save_with_retry(workbook)
        logger.info("Updated %s for tender: %s", sheet_name, reference)

    def _load_ready_workbook(self):
        if not self.workbook_path.exists():
            self._initialize_sync()

        workbook = load_workbook(self.workbook_path)
        self._rename_legacy_sheets(workbook)
        for sheet_name, headers in SHEET_HEADERS.items():
            worksheet = self._get_or_create_sheet(workbook, sheet_name)
            self._ensure_headers(worksheet, headers)
        return workbook

    @staticmethod
    def _rename_legacy_sheets(workbook) -> None:
        """Rename pre-milestone sheet names while staying in one workbook."""
        for old_name, new_name in LEGACY_SHEET_NAMES.items():
            if old_name in workbook.sheetnames and new_name not in workbook.sheetnames:
                workbook[old_name].title = new_name

    @staticmethod
    def _get_or_create_sheet(workbook, sheet_name: str) -> Worksheet:
        if sheet_name in workbook.sheetnames:
            return workbook[sheet_name]
        return workbook.create_sheet(sheet_name)

    @staticmethod
    def _ensure_headers(worksheet: Worksheet, expected_headers: list[str]) -> None:
        existing_headers = [
            worksheet.cell(row=1, column=column).value
            for column in range(1, worksheet.max_column + 1)
        ]

        if all(value is None for value in existing_headers):
            for column, header in enumerate(expected_headers, start=1):
                worksheet.cell(row=1, column=column, value=header)
            worksheet.freeze_panes = "A2"
            return

        if existing_headers == expected_headers:
            return

        header_indexes = {
            header: index
            for index, header in enumerate(existing_headers, start=1)
            if isinstance(header, str)
        }
        rows = []

        for row_number in range(2, worksheet.max_row + 1):
            rows.append(
                [
                    worksheet.cell(
                        row=row_number,
                        column=header_indexes.get(header, 0),
                    ).value
                    if header in header_indexes
                    else None
                    for header in expected_headers
                ]
            )

        worksheet.delete_rows(1, worksheet.max_row)

        for column, header in enumerate(expected_headers, start=1):
            worksheet.cell(row=1, column=column, value=header)

        for row_values in rows:
            worksheet.append(row_values)

        worksheet.freeze_panes = "A2"

    @staticmethod
    def _format_header(worksheet: Worksheet) -> None:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"

    @staticmethod
    def _header_indexes(worksheet: Worksheet) -> dict[str, int]:
        return {
            cell.value: cell.column
            for cell in worksheet[1]
            if isinstance(cell.value, str)
        }

    @staticmethod
    def _normalize_reference(reference: str) -> str:
        return reference.strip().casefold()

    def _reference_exists(self, worksheet: Worksheet, reference: str) -> bool:
        return self._find_reference_row(worksheet, reference) is not None

    def _find_reference_row(
        self,
        worksheet: Worksheet,
        reference: str,
    ) -> int | None:
        target = self._normalize_reference(reference)

        for row in worksheet.iter_rows(min_row=2, max_col=1):
            cell_value = row[0].value
            if cell_value is None:
                continue
            if self._normalize_reference(str(cell_value)) == target:
                return row[0].row

        return None

    @staticmethod
    def _raw_tender_row(tender: Tender) -> list[Any]:
        return [
            tender.reference,
            tender.tender_id,
            tender.title,
            tender.organisation,
            tender.department,
            tender.tender_number,
            tender.tender_category,
            tender.tender_type,
            tender.form_of_contract,
            tender.withdrawal_allowed,
            tender.no_of_covers,
            tender.payment_mode,
            tender.published_date,
            tender.doc_download_start_date,
            tender.doc_download_end_date,
            tender.clarification_start_date,
            tender.clarification_end_date,
            tender.bid_submission_start_date,
            tender.bid_submission_end_date,
            tender.closing_date,
            tender.opening_date,
            tender.pre_bid_meeting_date,
            tender.pre_bid_meeting_place,
            tender.bid_opening_place,
            tender.tender_value,
            tender.emd,
            tender.emd_exemption,
            tender.emd_fee_type,
            tender.tender_fee,
            tender.fee_payable_to,
            tender.fee_payable_at,
            tender.description,
            tender.product_category,
            tender.sub_category,
            tender.contract_type,
            tender.bid_validity_days,
            tender.period_of_work_days,
            tender.location,
            tender.pincode,
            tender.pre_qualification,
            tender.inviting_authority_name,
            tender.inviting_authority_address,
            tender.pdf_url,
            tender.pdf_local_path,
            ", ".join(tender.supporting_documents) if tender.supporting_documents else None,
        ]

    @staticmethod
    def _resolve_header(field_name: str, aliases: Mapping[str, str]) -> str:
        normalized = field_name.strip().lower().replace(" ", "_")
        header = aliases.get(normalized, field_name)

        valid_headers = set(AI_ANALYSIS_HEADERS + PROCESSING_STATUS_HEADERS)
        if header not in valid_headers:
            raise ValueError(f"Unsupported Excel field: {field_name}")

        return header

    @staticmethod
    def _cell_value(value: Any) -> Any:
        if isinstance(value, (list, tuple, set)):
            return ", ".join(str(item) for item in value)
        return value

    @staticmethod
    def _truthy(value: Any) -> bool:
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return value == 1
        if isinstance(value, str):
            return value.strip().casefold() in {"true", "yes", "1"}
        return False

    @staticmethod
    def _now() -> str:
        return datetime.now(timezone.utc).replace(microsecond=0).isoformat()
