"""
FastAPI server for the Tender Intelligence API.

Serves tender data from the Excel workbook for ENERGDive integration.
Run via: python app.py serve
"""

from __future__ import annotations

import asyncio
import math
from pathlib import Path
from typing import Optional

from fastapi import BackgroundTasks, FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from openpyxl import load_workbook

from api.schemas import (
    AIAnalysis,
    HealthResponse,
    ProcessingStatus,
    ScrapeJobResponse,
    SectorCount,
    StatsResponse,
    TenderBase,
    TenderListResponse,
    TenderResponse,
)
from config import settings
from utils.logger import get_logger

logger = get_logger("api.server")

app = FastAPI(
    title="Tender Intelligence API",
    description="ENERGDive Tender Intelligence — scrape, analyze, and serve government energy tenders.",
    version="1.0.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def _read_workbook() -> dict[str, list[dict[str, any]]]:
    """Read all sheets from the Excel workbook into dicts."""
    path = Path(settings.EXCEL_WORKBOOK_PATH)
    if not path.exists():
        return {"Raw_Tenders": [], "AI_Analysis": [], "Processing_Status": []}

    wb = load_workbook(path, read_only=True, data_only=True)
    result = {}

    for sheet_name in ["Raw_Tenders", "AI_Analysis", "Processing_Status"]:
        if sheet_name not in wb.sheetnames:
            result[sheet_name] = []
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            result[sheet_name] = []
            continue

        headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
        data = []
        for row in rows[1:]:
            if all(v is None for v in row):
                continue
            record = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            data.append(record)
        result[sheet_name] = data

    wb.close()
    return result


def _build_tender_response(ref: str, data: dict) -> TenderResponse:
    """Build a full TenderResponse from workbook data for a given reference."""
    raw_list = data.get("Raw_Tenders", [])
    ai_list = data.get("AI_Analysis", [])
    status_list = data.get("Processing_Status", [])

    raw = next((r for r in raw_list if str(r.get("Reference", "")).strip() == ref), None)
    ai = next((r for r in ai_list if str(r.get("Reference", "")).strip() == ref), None)
    status = next((r for r in status_list if str(r.get("Reference", "")).strip() == ref), None)

    tender_base = TenderBase(
        reference=ref,
        tender_id=raw.get("Tender ID") if raw else None,
        title=raw.get("Title") if raw else None,
        organisation=raw.get("Organisation") if raw else None,
        department=raw.get("Department") if raw else None,
        tender_number=raw.get("Tender Number") if raw else None,
        tender_category=raw.get("Tender Category") if raw else None,
        tender_type=raw.get("Tender Type") if raw else None,
        form_of_contract=raw.get("Form Of Contract") if raw else None,
        withdrawal_allowed=raw.get("Withdrawal Allowed") if raw else None,
        no_of_covers=str(raw.get("No. of Covers")) if raw and raw.get("No. of Covers") else None,
        payment_mode=raw.get("Payment Mode") if raw else None,
        tender_value=str(raw.get("Tender Value")) if raw and raw.get("Tender Value") else None,
        emd=str(raw.get("EMD")) if raw and raw.get("EMD") else None,
        emd_exemption=raw.get("EMD Exemption") if raw else None,
        emd_fee_type=raw.get("EMD Fee Type") if raw else None,
        tender_fee=str(raw.get("Tender Fee")) if raw and raw.get("Tender Fee") else None,
        fee_payable_to=raw.get("Fee Payable To") if raw else None,
        fee_payable_at=raw.get("Fee Payable At") if raw else None,
        description=raw.get("Description") if raw else None,
        product_category=raw.get("Product Category") if raw else None,
        sub_category=raw.get("Sub category") if raw else None,
        contract_type=raw.get("Contract Type") if raw else None,
        bid_validity_days=str(raw.get("Bid Validity")) if raw and raw.get("Bid Validity") else None,
        period_of_work_days=str(raw.get("Period Of Work")) if raw and raw.get("Period Of Work") else None,
        location=raw.get("Location") if raw else None,
        pincode=str(raw.get("Pincode")) if raw and raw.get("Pincode") else None,
        pre_qualification=raw.get("Pre Qualification") if raw else None,
        published_date=str(raw.get("Published Date")) if raw and raw.get("Published Date") else None,
        doc_download_start_date=str(raw.get("Document Download Start Date")) if raw and raw.get("Document Download Start Date") else None,
        doc_download_end_date=str(raw.get("Document Download End Date")) if raw and raw.get("Document Download End Date") else None,
        clarification_start_date=str(raw.get("Clarification Start Date")) if raw and raw.get("Clarification Start Date") else None,
        clarification_end_date=str(raw.get("Clarification End Date")) if raw and raw.get("Clarification End Date") else None,
        bid_submission_start_date=str(raw.get("Bid Submission Start Date")) if raw and raw.get("Bid Submission Start Date") else None,
        bid_submission_end_date=str(raw.get("Bid Submission End Date")) if raw and raw.get("Bid Submission End Date") else None,
        closing_date=str(raw.get("Closing Date")) if raw and raw.get("Closing Date") else None,
        opening_date=str(raw.get("Opening Date")) if raw and raw.get("Opening Date") else None,
        pre_bid_meeting_date=str(raw.get("Pre Bid Meeting Date")) if raw and raw.get("Pre Bid Meeting Date") else None,
        pre_bid_meeting_place=raw.get("Pre Bid Meeting Place") if raw else None,
        bid_opening_place=raw.get("Bid Opening Place") if raw else None,
        inviting_authority_name=raw.get("Inviting Authority Name") if raw else None,
        inviting_authority_address=raw.get("Inviting Authority Address") if raw else None,
        pdf_url=raw.get("PDF URL") if raw else None,
        pdf_path=raw.get("PDF Path") if raw else None,
        supporting_documents=_to_list(raw.get("Supporting Documents")) if raw else None,
    )

    ai_analysis = None
    if ai:
        keywords_raw = ai.get("Keywords")
        tags_raw = ai.get("Tags")
        companies_raw = ai.get("Companies Interested")

        ai_analysis = AIAnalysis(
            is_energy_related=_to_bool(ai.get("Energy Related")),
            energy_sector=ai.get("Energy Sector"),
            headline=ai.get("Headline"),
            summary=ai.get("Summary"),
            seo_title=ai.get("SEO Title"),
            meta_description=ai.get("Meta Description"),
            slug=ai.get("Slug"),
            keywords=_to_list(keywords_raw),
            tags=_to_list(tags_raw),
            importance_score=_to_int(ai.get("Importance Score")),
            publish_recommendation=ai.get("Publish Recommendation"),
            why_it_matters=ai.get("Why It Matters"),
            industry_impact=ai.get("Industry Impact"),
            companies_interested=_to_list(companies_raw),
            eligibility=ai.get("Eligibility"),
            required_documents=ai.get("Required Documents"),
            important_dates=ai.get("Important Dates"),
            risk_level=ai.get("Risk Level"),
        )

    processing_status = None
    if status:
        processing_status = ProcessingStatus(
            downloaded=_to_bool(status.get("Downloaded")),
            pdf_parsed=_to_bool(status.get("PDF Parsed")),
            ai_completed=_to_bool(status.get("AI Completed")),
            created_at=str(status.get("Created At")) if status.get("Created At") else None,
            updated_at=str(status.get("Updated At")) if status.get("Updated At") else None,
        )

    return TenderResponse(
        tender=tender_base,
        ai_analysis=ai_analysis,
        processing_status=processing_status,
    )


def _to_bool(value) -> Optional[bool]:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value == 1
    if isinstance(value, str):
        return value.strip().lower() in {"true", "yes", "1"}
    return None


def _to_int(value) -> Optional[int]:
    if value is None:
        return None
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def _to_list(value) -> Optional[list[str]]:
    if value is None:
        return None
    if isinstance(value, str):
        return [s.strip() for s in value.split(",") if s.strip()]
    if isinstance(value, list):
        return value
    return None


@app.get("/health", response_model=HealthResponse)
async def health_check():
    path = Path(settings.EXCEL_WORKBOOK_PATH)
    return HealthResponse(
        status="healthy",
        version="1.0.0",
        excel_path=str(path.resolve()),
        excel_exists=path.exists(),
    )


@app.get("/api/v1/tenders", response_model=TenderListResponse)
async def list_tenders(
    page: int = Query(1, ge=1, description="Page number"),
    page_size: int = Query(20, ge=1, le=100, description="Items per page"),
    energy_only: bool = Query(False, description="Filter to energy-related tenders only"),
    sector: Optional[str] = Query(None, description="Filter by energy sector"),
    search: Optional[str] = Query(None, description="Search in title, org, description"),
):
    data = _read_workbook()
    raw_tenders = data.get("Raw_Tenders", [])
    ai_data = {str(r.get("Reference", "")).strip(): r for r in data.get("AI_Analysis", [])}

    references = [str(r.get("Reference", "")).strip() for r in raw_tenders if r.get("Reference")]

    # Apply filters
    filtered_refs = []
    for ref in references:
        ai = ai_data.get(ref, {})

        if energy_only and not _to_bool(ai.get("Energy Related")):
            continue

        if sector and ai.get("Energy Sector") != sector:
            continue

        if search:
            raw = next((r for r in raw_tenders if str(r.get("Reference", "")).strip() == ref), {})
            searchable = " ".join([
                str(raw.get("Title", "")),
                str(raw.get("Organisation", "")),
                str(raw.get("Description", "")),
            ]).lower()
            if search.lower() not in searchable:
                continue

        filtered_refs.append(ref)

    total = len(filtered_refs)
    total_pages = max(1, math.ceil(total / page_size))
    start = (page - 1) * page_size
    end = start + page_size
    page_refs = filtered_refs[start:end]

    tenders = [_build_tender_response(ref, data) for ref in page_refs]

    return TenderListResponse(
        tenders=tenders,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )


@app.get("/api/v1/tenders/energy", response_model=TenderListResponse)
async def list_energy_tenders(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    sector: Optional[str] = Query(None, description="Filter by energy sector"),
):
    """Primary ENERGDive endpoint — returns only energy-related tenders."""
    return await list_tenders(
        page=page,
        page_size=page_size,
        energy_only=True,
        sector=sector,
        search=None,
    )


@app.get("/api/v1/tenders/{reference:path}", response_model=TenderResponse)
async def get_tender(reference: str):
    data = _read_workbook()
    raw_refs = [str(r.get("Reference", "")).strip() for r in data.get("Raw_Tenders", [])]

    if reference not in raw_refs:
        raise HTTPException(status_code=404, detail=f"Tender '{reference}' not found.")

    return _build_tender_response(reference, data)


@app.get("/api/v1/stats", response_model=StatsResponse)
async def get_stats():
    data = _read_workbook()
    raw_tenders = data.get("Raw_Tenders", [])
    ai_data = data.get("AI_Analysis", [])
    status_data = data.get("Processing_Status", [])

    total = len(raw_tenders)
    energy_related = sum(1 for r in ai_data if _to_bool(r.get("Energy Related")))
    non_energy = total - energy_related
    ai_completed = sum(1 for r in status_data if _to_bool(r.get("AI Completed")))
    ai_pending = total - ai_completed
    pdfs_downloaded = sum(1 for r in status_data if _to_bool(r.get("Downloaded")))

    sectors: dict[str, int] = {}
    for row in ai_data:
        s = row.get("Energy Sector")
        if s and isinstance(s, str) and s.strip():
            sectors[s.strip()] = sectors.get(s.strip(), 0) + 1

    sector_breakdown = [
        SectorCount(sector=k, count=v)
        for k, v in sorted(sectors.items(), key=lambda x: -x[1])
    ]

    return StatsResponse(
        total_tenders=total,
        energy_related=energy_related,
        non_energy=non_energy,
        ai_completed=ai_completed,
        ai_pending=ai_pending,
        pdfs_downloaded=pdfs_downloaded,
        sector_breakdown=sector_breakdown,
    )


@app.get("/api/v1/export/excel")
async def export_excel():
    path = Path(settings.EXCEL_WORKBOOK_PATH)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Excel workbook not found.")
    return FileResponse(
        path=str(path.resolve()),
        filename="tenders.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


async def _run_scrape_pipeline() -> ScrapeJobResponse:
    """Run the full scrape + AI analysis pipeline."""
    from services.ai_service import AIConfigurationError, AIService, AIServiceError, AIResponseValidationError
    from services.excel_service import ExcelService
    from scrapers.cppp_scraper import CPPPScraper

    storage = ExcelService(settings.EXCEL_WORKBOOK_PATH)
    await storage.initialize()

    try:
        ai_service = AIService()
    except AIConfigurationError as exc:
        logger.error("AI configuration failed: %s", exc)
        ai_service = None

    async def should_process(tender):
        if not await storage.tender_exists(tender.reference):
            return True
        if not await storage.ai_analysis_completed(tender.reference):
            return True
        return False

    scraper = CPPPScraper()
    tenders = await scraper.fetch(should_process=should_process)

    saved = 0
    ai_ok = 0
    ai_fail = 0

    for tender in tenders:
        if await storage.append_tender(tender):
            saved += 1

        try:
            await storage.update_processing_status(
                tender.reference,
                {"downloaded": bool(tender.pdf_local_path), "pdf_parsed": bool(tender.extracted_text)},
            )
        except Exception as exc:
            logger.exception("Status update failed for %s: %s", tender.reference, exc)

        if ai_service:
            try:
                analysis = await ai_service.analyze(tender)
                await storage.update_ai_analysis(tender.reference, analysis.to_excel_fields())
                await storage.update_processing_status(tender.reference, {"ai_completed": True})
                ai_ok += 1
            except Exception as exc:
                logger.error("AI failed for %s: %s", tender.reference, exc)
                ai_fail += 1
        else:
            ai_fail += 1

    return ScrapeJobResponse(
        status="completed",
        message=f"Scrape pipeline finished. {len(tenders)} tenders processed.",
        tenders_found=len(tenders),
        tenders_saved=saved,
        ai_completed=ai_ok,
        ai_failed=ai_fail,
    )


# Store for background job result
_scrape_result: dict = {"running": False, "last_result": None}


async def _background_scrape():
    global _scrape_result
    _scrape_result["running"] = True
    try:
        result = await _run_scrape_pipeline()
        _scrape_result["last_result"] = result.model_dump()
    except Exception as exc:
        logger.exception("Background scrape failed: %s", exc)
        _scrape_result["last_result"] = {"status": "failed", "message": str(exc)}
    finally:
        _scrape_result["running"] = False


@app.post("/api/v1/scrape", response_model=ScrapeJobResponse)
async def trigger_scrape(background_tasks: BackgroundTasks):
    if _scrape_result["running"]:
        return ScrapeJobResponse(
            status="already_running",
            message="A scrape job is already in progress.",
        )

    background_tasks.add_task(_background_scrape)
    return ScrapeJobResponse(
        status="started",
        message="Scrape pipeline started in the background. Check /api/v1/stats for progress.",
    )
